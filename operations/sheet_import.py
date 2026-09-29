"""Import items from a spreadsheet-style file: CSV, TSV, Excel (.xlsx) or JSON.

Used by the System page's "Import database" button for anything that isn't a
SQLite database. Rows are matched to items by SKU: a new SKU adds an item, a
known SKU updates it, and items not in the file are left alone. A blank cell
never erases a value that is already there.

Column names are matched loosely, so Dispodex's own exports, the partner
exports and most hand-made sheets work ("Item SKU", "sku", "SKU #" …). Nothing
is lost: a value that doesn't fit its field (an unknown condition, a price
that isn't a number, an extra column) is written into the item's notes with
its column name. Every change goes through the normal save, so it shows in the
item's History under the importer's name and is queued for Square.
"""
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from django.db import transaction

from core.skus import normalize_sku
from inventory.models import (
    LEGACY_CONDITION_MAP,
    CompatibleOS,
    Condition,
    Functional,
    Item,
    Status,
    YesNo,
    coerce_status,
)
from inventory.services import items as item_service
from inventory.services.exports import EXPORT_COLUMNS, PARTNER_COLUMNS
from operations import excel_import

SHEET_SUFFIXES = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".json"}
# Formats people often have that we can't read directly: say how to convert them.
CONVERT_SUFFIXES = {".xls", ".ods", ".numbers", ".xlsb"}


class SheetError(Exception):
    """The file can't be imported; the message is shown to the person."""


def _norm(text) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(text or "").lower())


# Loose column name -> Item field.
SYSTEM_FIELDS = {"id", "created_at", "updated_at"}  # set by the database, never from a file
COLUMN_ALIASES: dict[str, str] = {}
for _field, _header in (*EXPORT_COLUMNS, *PARTNER_COLUMNS):
    if _field in SYSTEM_FIELDS:
        continue
    COLUMN_ALIASES.setdefault(_norm(_header), _field)
    COLUMN_ALIASES.setdefault(_norm(_field), _field)
COLUMN_ALIASES.update({
    _norm(alias): name for name, aliases in {
        "sku": ["Item SKU", "SKU #", "SKU Number", "Stock Code", "Product SKU", "Inventory SKU"],
        "what_is_it": ["Item", "Item Name", "Name", "Title", "Description", "Product", "Product Name"],
        "brand_model": ["Brand", "Model", "Brand & Model", "Brand and Model", "Model Number"],
        "quantity": ["Quantity", "Qty On Hand", "Count", "Stock"],
        "price": ["Price ($)", "Sale Price", "List Price", "eBay Price", "Asking Price"],
        "where_it_goes": ["Location", "Where", "Shelf", "Bin"],
        "ebay_category": ["Category"],
        "serial_number": ["Serial", "Serial #", "S/N", "SN"],
        "ssd_gb": ["Storage", "SSD", "Hard Drive", "HDD", "Disk"],
        "graphics_card": ["Graphics", "GPU", "Video Card"],
        "power_on": ["Powers On", "Power On?"],
        "cords_adapters": ["Cords/Adapters Included", "Cords", "Adapters", "Charger"],
        "os": ["Operating System"],
        "date_received": ["Received", "Date In", "Intake Date"],
        "notes": ["Note", "Comments", "Comment", "Memo"],
        "status": ["Lane", "Stage"],
    }.items() for alias in aliases
})
# System columns from our own exports: the database sets these itself.
IGNORED = {_norm(name) for name in ("ID", "Created", "Updated", "created_at", "updated_at", "Photos", "Photo")}

CHOICES = {
    "condition": Condition, "functional": Functional, "power_on": YesNo, "cords_adapters": YesNo,
    "keep_items_together": YesNo, "picture_taken": YesNo, "compatible_os": CompatibleOS,
}
FLAGS = {"diagnostics_test_ran", "wifi_card_installed"}
TRUE_WORDS = {"yes", "y", "true", "1", "x", "on"}
FALSE_WORDS = {"no", "n", "false", "0", "off"}


@dataclass
class SheetReport:
    rows: int = 0
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    skipped: list[str] = field(default_factory=list)
    unknown_columns: list[str] = field(default_factory=list)
    kept_in_notes: int = 0

    def lines(self) -> list[str]:
        if not self.created and not self.updated and self.unchanged:
            # Common when re-importing an export of this same Dispodex: say so plainly,
            # otherwise "0 new, 0 updated" reads like the file was rejected.
            lines = [f"Nothing needed changing: all {self.unchanged} rows in this file already match what's in Dispodex."]
        else:
            lines = [f"{self.rows} rows read: {self.created} new items, {self.updated} updated, {self.unchanged} already up to date."]
        if self.skipped:
            lines.append(f"{len(self.skipped)} rows skipped: " + "; ".join(self.skipped[:5]) + (" …" if len(self.skipped) > 5 else ""))
        if self.unknown_columns:
            lines.append("Columns kept in notes (no matching field): " + ", ".join(self.unknown_columns) + ".")
        if self.kept_in_notes:
            lines.append(f"{self.kept_in_notes} values didn't fit their field and were kept in the item's notes.")
        return lines


# ── reading files ────────────────────────────────────────────────────────────
def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _table(header: list, body) -> tuple[list[str], list[dict]]:
    headers = [_text(h) for h in header]
    rows = []
    for values in body:
        row = {h: _text(v) for h, v in zip(headers, values) if h}
        if any(row.values()):
            rows.append(row)
    return [h for h in headers if h], rows


def _read_delimited(path: Path) -> tuple[list[str], list[dict]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        raise SheetError("That file is empty.")
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines[:20]), delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    records = list(csv.reader(io.StringIO("\n".join(lines)), dialect))
    return _table(records[0], records[1:])


def _openpyxl_records(path: Path) -> list[list]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(r) for r in book.worksheets[0].iter_rows(values_only=True)]
    finally:
        book.close()


def is_old_app_export(path: Path) -> bool:
    """True for the old PHP app's "Excel with photos" export, which openpyxl can't read
    (its columns after Z are named "[", "\\" …) but ``excel_import`` can, photos included."""
    try:
        _openpyxl_records(path)
        return False
    except Exception:  # noqa: BLE001 — any openpyxl failure: try the old-export reader
        try:
            excel_import.read_workbook(path)[2].close()
            return True
        except excel_import.ExcelImportError:
            return False


def _read_xlsx(path: Path) -> tuple[list[str], list[dict]]:
    try:
        records = _openpyxl_records(path)
    except Exception:  # noqa: BLE001 — the old app's export fails partway through reading
        try:
            headers, rows, book = excel_import.read_workbook(path)
        except excel_import.ExcelImportError as exc:
            raise SheetError(f"That Excel file can't be read: {exc}") from exc
        book.close()
        return headers, [row.values for row in rows]
    while records and not any(_text(v) for v in records[0]):
        records.pop(0)
    if not records:
        raise SheetError("The first sheet of that workbook is empty.")
    return _table(records[0], records[1:])


def _read_json(path: Path) -> tuple[list[str], list[dict]]:
    try:
        data = json.loads(path.read_bytes().decode("utf-8-sig"))
    except (ValueError, UnicodeDecodeError) as exc:
        raise SheetError(f"That JSON file can't be read: {exc}") from exc
    if isinstance(data, dict):
        data = next((v for k, v in data.items() if isinstance(v, list) and k.lower() in ("items", "rows", "data", "records")), None)
    if not isinstance(data, list) or not all(isinstance(row, dict) for row in data):
        raise SheetError('A JSON import must be a list of items, e.g. [{"sku": "A-1", "what_is_it": "Laptop"}].')
    headers: list[str] = []
    for row in data:
        headers += [k for k in row if k not in headers]
    rows = [{k: _text(v) for k, v in row.items()} for row in data]
    return headers, [r for r in rows if any(r.values())]


def read(path: Path, original_name: str) -> tuple[list[str], list[dict]]:
    """Headers and rows (header -> text). Raises SheetError without changing anything."""
    suffix = Path(original_name).suffix.lower()
    if suffix in CONVERT_SUFFIXES:
        raise SheetError(f"{suffix} files can't be read directly. Open it and use Save As → Excel Workbook (.xlsx) or CSV, then import that.")
    if suffix not in SHEET_SUFFIXES:
        raise SheetError("That type of file can't be imported. Use a database (.sqlite3, .sqlite, .db), "
                         "a spreadsheet (.xlsx, .csv, .tsv) or a .json file.")
    if suffix in (".xlsx", ".xlsm"):
        headers, rows = _read_xlsx(path)
    elif suffix == ".json":
        headers, rows = _read_json(path)
    else:
        headers, rows = _read_delimited(path)
    if not any(COLUMN_ALIASES.get(_norm(h)) == "sku" for h in headers):
        raise SheetError('The file needs a "SKU" column so each row can be matched to an item. '
                         "Found: " + (", ".join(headers[:12]) or "no column names") + ".")
    if not rows:
        raise SheetError("The file has column names but no rows.")
    return headers, rows


# ── turning cells into field values ─────────────────────────────────────────
def _parse_date(text: str) -> date | None:
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%m-%d-%Y", "%d %b %Y", "%b %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        return None


def _choice(field_name: str, text: str) -> str | None:
    key = _norm(text)
    if field_name == "condition" and text.strip().lower() in LEGACY_CONDITION_MAP:
        return LEGACY_CONDITION_MAP[text.strip().lower()]
    choices = CHOICES[field_name]
    if choices is YesNo or choices is Functional:
        if key in TRUE_WORDS:
            return "Yes"
        if key in FALSE_WORDS:
            return "No"
    for value, label in choices.choices:
        if key in (_norm(value), _norm(label)) or (field_name == "compatible_os" and key == _norm(label).replace("windows", "win")):
            return value
    return None


def _status(text: str) -> str | None:
    key = _norm(text)
    for value, label in Status.choices:
        if key in (_norm(value), _norm(label)):
            return value
    coerced = coerce_status(text)
    return coerced if coerced != Status.INTAKE or key == "intake" else None


def convert(field_name: str, text: str):
    """The value for ``field_name``, or raise ValueError when it doesn't fit."""
    if field_name == "sku":
        return text.strip().upper()
    if field_name == "status":
        value = _status(text)
    elif field_name in CHOICES:
        value = _choice(field_name, text)
    elif field_name in FLAGS:
        value = True if _norm(text) in TRUE_WORDS else False if _norm(text) in FALSE_WORDS else None
    elif field_name == "price":
        try:
            value = item_service.parse_price(text)
        except item_service.ItemError:
            value = None
    elif field_name == "quantity":
        try:
            value = int(Decimal(text.replace(",", "")))
        except ArithmeticError:
            value = None
        if value is not None and value < 1:
            value = None
    elif field_name == "date_received":
        value = _parse_date(text)
    else:
        limit = Item._meta.get_field(field_name).max_length
        value = text if not limit or len(text) <= limit else None
    if value is None:
        raise ValueError(text)
    return value


# ── importing ────────────────────────────────────────────────────────────────
def apply(headers: list[str], rows: list[dict], *, source: str) -> SheetReport:
    """Add or update one item per row. All or nothing: any crash undoes every row."""
    report = SheetReport(rows=len(rows))
    mapping: dict[str, str] = {}
    for header in headers:
        target = COLUMN_ALIASES.get(_norm(header))
        if target and target not in mapping.values():
            mapping[header] = target
        elif _norm(header) not in IGNORED:
            report.unknown_columns.append(header)
    sku_header = next(h for h, f in mapping.items() if f == "sku")

    with transaction.atomic():
        for number, row in enumerate(rows, start=2):
            sku = (row.get(sku_header) or "").strip()
            if not normalize_sku(sku):
                report.skipped.append(f"row {number} has no SKU")
                continue
            values, extras = {}, []
            for header, field_name in mapping.items():
                text = (row.get(header) or "").strip()
                if not text or field_name == "sku":
                    continue
                try:
                    values[field_name] = convert(field_name, text)
                except ValueError:
                    extras.append(f"{header}: {text}")
            extras += [f"{h}: {row[h].strip()}" for h in report.unknown_columns if (row.get(h) or "").strip()]

            item = item_service.find_item(sku)
            notes = values.get("notes", item.notes if item else "")
            new_lines = [line for line in extras if line not in notes]
            if new_lines:
                report.kept_in_notes += len(new_lines)
                values["notes"] = (notes + "\n" if notes else "") + f"[From {source}]\n" + "\n".join(new_lines)
            if item and all(getattr(item, f) == v for f, v in values.items()):
                report.unchanged += 1
                continue
            try:
                result = item_service.save_intake({"sku": sku, **values})
            except item_service.ItemError as exc:
                report.skipped.append(f"row {number} ({sku}): {exc}")
                continue
            if result.created:
                report.created += 1
            else:
                report.updated += 1
    return report
