import csv
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from django.urls import reverse

from archive.importer import import_csv
from archive.models import ArchiveItem
from inventory.models import Item, ScriptCache
from inventory.services import labels, scripts
from inventory.services.search import palette_search, suggestions

pytestmark = pytest.mark.django_db
FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def stock(db):
    Item.objects.create(sku="LAP-1", what_is_it="Laptop", brand_model="Dell Latitude", price="100", quantity=2)
    Item.objects.create(sku="LAP-2", what_is_it="Laptop", brand_model="HP EliteBook", price="300", status="sold")
    Item.objects.create(sku="DESK-1", what_is_it="Desktop", serial_number="SN12345", price=None)


@pytest.mark.parametrize("url", [
    "/", "/intake/?new=1", "/intake/?sku=LAP-1", "/lookup/", "/board/", "/archive/", "/scripts/",
    "/scripts/?sku=LAP-1", "/system/", "/login/",
])
def test_pages_render(client, stock, url):
    page = client.get(url)
    assert page.status_code == 200
    assert b"Listing images" not in page.content  # hidden until it's finished


def test_listing_images_is_off_until_finished(client, stock):
    assert client.get("/listing-images/").status_code == 404
    assert client.post("/api/listing-images/upload/", {"sku": "LAP-1"}).status_code == 404
    assert client.post("/api/listing-images/LAP-1/layout/", "{}", content_type="application/json").status_code == 404


@pytest.mark.parametrize("url", ["/listing-images/", "/listing-images/?sku=LAP-1"])
def test_listing_images_comes_back_with_the_switch(client, stock, settings, url):
    settings.PINKSHEET = {**settings.PINKSHEET, "LISTING_IMAGES": True}
    assert client.get(url).status_code == 200
    assert b"Listing images" in client.get("/intake/?sku=LAP-1").content


def test_lookup_filters(client, stock):
    page = client.get("/lookup/?q=lap").content.decode()
    assert "LAP-1" in page and "LAP-2" in page and "DESK-1" not in page
    page = client.get("/lookup/?status=sold").content.decode()
    assert "LAP-2" in page and "LAP-1" not in page
    page = client.get("/lookup/?min_price=150").content.decode()
    assert "LAP-2" in page and "LAP-1" not in page
    page = client.get("/lookup/?gap=no-price").content.decode()
    assert "DESK-1" in page and "LAP-1" not in page
    page = client.get("/lookup/?scope=active").content.decode()
    assert "LAP-2" not in page


def test_lookup_total_value_uses_quantity(client, stock):
    assert "$500.00" in client.get("/lookup/").content.decode()  # 100×2 + 300


def test_suggestions_and_palette(stock):
    assert [s["value"] for s in suggestions("lap")] == ["LAP-2", "LAP-1"]
    assert palette_search("sn12345")[0]["sku"] == "DESK-1"
    assert palette_search("LAP-1")[0]["sku"] == "LAP-1"


def test_board_cards_grouped_by_lane(client, stock):
    lanes = client.get(reverse("api_board_cards")).json()["lanes"]
    assert {c["sku"] for c in lanes["intake"]} == {"LAP-1", "DESK-1"}
    assert [c["sku"] for c in lanes["sold"]] == ["LAP-2"]


def test_card_redirect_depends_on_device(client, stock):
    phone = client.get("/card/LAP-1/", HTTP_USER_AGENT="Mozilla/5.0 (iPhone; CPU iPhone OS 17_0)")
    assert phone["Location"] == "/m/LAP-1/"
    desktop = client.get("/card/LAP-1/", HTTP_USER_AGENT="Mozilla/5.0 (Windows NT 10.0)")
    assert desktop["Location"] == "/intake/?sku=LAP-1"


def test_old_php_links_redirect(client, stock):
    assert client.get("/card.php?sku=LAP-1")["Location"] == "/card/LAP-1/"
    assert client.get("/intake.php?sku=LAP-1")["Location"] == "/intake/?sku=LAP-1"
    assert client.get("/intake.php?clear_draft=1")["Location"] == "/intake/?new=1"
    assert client.get("/kanban.php")["Location"] == "/board/"
    assert client.get("/photo.php?id=5&thumb=1")["Location"] == "/photos/5/?thumb=1"


def test_mobile_and_print_pages(client, stock):
    assert b"Mark sold" in client.get("/m/LAP-1/").content
    assert client.get("/m/NOPE/").status_code == 404
    assert b"Dell Latitude" in client.get("/print/LAP-1/").content


# ── exports ──────────────────────────────────────────────────────────────────
def test_csv_export(client, stock):
    response = client.get(reverse("export_csv") + "?scope=active")
    text = response.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0][0] == "SKU"
    assert {r[0] for r in rows[1:]} == {"LAP-1", "DESK-1"}
    assert "inventory_active_" in response["Content-Disposition"]


def test_zip_export_has_folders_and_photos(client, stock, image):
    client.post(reverse("api_photo_upload"), {"sku": "LAP-1", "photo": image(name="front.jpg")})
    response = client.get(reverse("export_zip"))
    archive = zipfile.ZipFile(io.BytesIO(b"".join(response.streaming_content)))
    names = archive.namelist()
    assert "LAP-1/01_front.png" in names
    assert "LAP-1/info.txt" in names
    assert any(n.startswith("inventory_") and n.endswith(".csv") for n in names)
    assert "Photos: 1" in archive.read("LAP-1/info.txt").decode()


def test_xlsx_export_embeds_photos(client, stock, image):
    client.post(reverse("api_photo_upload"), {"sku": "LAP-1", "photo": image()})
    response = client.get(reverse("export_xlsx"))
    book = openpyxl.load_workbook(io.BytesIO(b"".join(response.streaming_content)))
    sheet = book.active
    assert sheet["A1"].value == "SKU"
    assert len(sheet._images) == 1
    skus = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
    assert set(skus) == {"LAP-1", "LAP-2", "DESK-1"}


# ── archive ──────────────────────────────────────────────────────────────────
def test_archive_csv_import_is_idempotent(tmp_path):
    path = tmp_path / "old.csv"
    path.write_text(
        "Inventory ID,Item SKU,Description,Sold Price,Sold Date,Created At\n"
        "41,iz-0105-55,Dell Optiplex 7040,$125.00,2025-08-29,2025-08-29 11:49:19.749 -0500\n"
        "42,IZ-0105-56,Latitude,,,\n",
        encoding="utf-8",
    )
    report = import_csv(path, source="Legacy", table="inventory")
    assert report.inserted == 2
    again = import_csv(path, source="Legacy", table="inventory")
    assert again.inserted == 0 and again.skipped == 2
    first = ArchiveItem.objects.get(legacy_id="41")
    assert first.sku_normalized == "IZ-0105-55"
    assert str(first.sold_price) == "125.00"
    assert first.sold_at.isoformat() == "2025-08-29"
    assert json.loads(first.legacy_payload)["Item SKU"] == "iz-0105-55"


def test_archive_search_and_export(client):
    ArchiveItem.objects.create(sku="OLD-1", title="Dell Optiplex", buyer="Jane", status="Sold", legacy_payload="{}")
    ArchiveItem.objects.create(sku="OLD-2", title="HP Printer", status="Archived", legacy_payload="{}")
    page = client.get("/archive/?q=optiplex").content.decode()
    assert "OLD-1" in page and "OLD-2" not in page
    rows = list(csv.reader(io.StringIO(client.get("/exports/archive.csv?status=archived").content.decode("utf-8-sig"))))
    assert [r[0] for r in rows[1:]] == ["OLD-2"]


# ── script builder & labels ──────────────────────────────────────────────────
@pytest.fixture
def shop_notes(settings, tmp_path):
    notes = tmp_path / "ebay_boilerplate.txt"
    notes.write_text("Please Read This First\nShips Mon-Fri\n", encoding="utf-8")
    settings.PINKSHEET = {**settings.PINKSHEET, "EBAY_BOILERPLATE_FILE": notes}
    return "Please Read This First\nShips Mon-Fri"


def _listing_answer(notes):
    """An answer to the prompt, with ChatGPT's (slightly reworded) copy of the shop notes."""
    return f"""Sure! Here's your listing:

**Recommended eBay Title:** Panasonic Lumix DMC-FS3 8.1MP Digital Camera

{notes.replace("Ships Mon-Fri", "Ships Monday-Friday")}

**Product Details**
- Brand/Model: Panasonic Lumix DMC-FS3
- Battery Health: Holds charge
- Inventory Number: CB2-1

**Suggested eBay Price:** $45-$60
**Recommended Shipping:** USPS Ground Advantage, about $6"""


def test_prompt_is_the_old_builders_full_listing_prompt(stock, shop_notes):
    prompt = scripts.build_prompt("LAP-1", Item.objects.get(sku="LAP-1"))
    assert prompt.startswith("Generate a concise eBay listing for this item.")
    assert f"include this EXACT boilerplate text (copy it exactly as shown):\n\n{shop_notes}\n\n3. THEN" in prompt
    assert "PROVIDED SPECS:\nBrand / Model: Dell Latitude" in prompt and "SKU: LAP-1" in prompt
    assert "End the Product Details section with: Inventory Number: LAP-1" in prompt
    assert "Suggested eBay price" in prompt and "No condition mentions" in prompt


def test_listing_answer_becomes_title_description_and_staff_notes(client, shop_notes):
    built = client.post(reverse("api_script_build"), {"chatgpt_text": _listing_answer(shop_notes)},
                        content_type="application/json").json()
    assert built["title"] == "Panasonic Lumix DMC-FS3 8.1MP Digital Camera" and built["title_length"] == 44
    # The shop notes exactly as written (not ChatGPT's reworded copy), then the Product Details.
    assert built["final_text"] == (
        "Please Read This First\nShips Mon-Fri\n\nProduct Details\n- Brand/Model: Panasonic Lumix DMC-FS3\n"
        "- Battery Health: Holds charge\n- Inventory Number: CB2-1"
    )
    assert built["staff_notes"] == "Suggested eBay Price: $45-$60\nRecommended Shipping: USPS Ground Advantage, about $6"


def test_answers_without_the_shop_notes_get_them_added(shop_notes):
    assert scripts.build_final_script("  Great laptop  ") == f"{shop_notes}\n\nGreat laptop"
    parsed = scripts.parse_answer('Title: Dell Latitude 5590 15.6"\n\nPowers on.')
    assert parsed.title == 'Dell Latitude 5590 15.6"' and parsed.description == f"{shop_notes}\n\nPowers on."
    assert scripts.build_final_script("") == "Paste the ChatGPT output first."


def test_script_cache_api(client, stock):
    loaded = client.get(reverse("api_script", args=["LAP-1"])).json()
    assert loaded["found"] and "LAP-1" in loaded["prompt_text"]
    client.post(reverse("api_script", args=["LAP-1"]),
                {"prompt_text": "p", "chatgpt_text": "c", "final_text": "f"}, content_type="application/json")
    cache = ScriptCache.objects.get()
    assert (cache.prompt_text, cache.chatgpt_text, cache.final_text, cache.state) == ("p", "c", "f", "ready")
    fresh = client.get(reverse("api_script", args=["LAP-1"]) + "?fresh=1").json()
    assert fresh["prompt_text"].startswith("Generate a concise eBay listing")


def test_saved_finals_from_before_are_rebuilt_but_edits_are_kept(client, stock, shop_notes):
    answer = _listing_answer(shop_notes)
    old_final = f"{scripts.final_boilerplate()}\n\n{answer}"
    ScriptCache.objects.create(sku_normalized="LAP-1", sku_display="LAP-1", chatgpt_text=answer, final_text=old_final)
    loaded = client.get(reverse("api_script", args=["LAP-1"])).json()
    assert "Sure!" not in loaded["final_text"] and loaded["final_text"].endswith("Inventory Number: CB2-1")
    assert loaded["title"].startswith("Panasonic") and loaded["staff_notes"].startswith("Suggested eBay Price")

    ScriptCache.objects.filter(sku_normalized="LAP-1").update(final_text="Hand-edited listing")
    assert client.get(reverse("api_script", args=["LAP-1"])).json()["final_text"] == "Hand-edited listing"


def test_builder_warns_when_the_shop_notes_file_is_missing(client, settings, tmp_path):
    settings.PINKSHEET = {**settings.PINKSHEET, "EBAY_BOILERPLATE_FILE": tmp_path / "missing.txt"}
    assert "listing notes file wasn&#x27;t found" in client.get("/scripts/").content.decode().replace("'", "&#x27;")
    (tmp_path / "missing.txt").write_text("Our policies\n", encoding="utf-8")
    assert "listing notes file" not in client.get("/scripts/").content.decode()


def test_zpl_matches_the_old_php_app_exactly():
    """Fixture generated by running the PHP app's zpl.php — labels must print identically."""
    cases = json.loads((FIXTURES / "zpl_cases.json").read_text(encoding="utf-8"))
    for case in cases:
        assert labels.generate_zpl(case["input"]) == case["zpl"], case["input"]


def test_label_endpoint(client, stock):
    data = client.get(reverse("api_label_zpl") + "?sku=LAP-1&preset=detail").json()
    assert data["ok"] and data["zpl"].startswith("^XA") and "^FDLA,LAP-1^FS" in data["zpl"]
    assert data["labelPreset"] == "Intake label"


def test_ebay_categories_endpoint(client):
    data = client.get(reverse("api_ebay_categories")).json()
    assert data["source"] == "bundled" and data["count"] > 50


# ── health & settings ────────────────────────────────────────────────────────
def test_health_endpoint(client, stock):
    data = client.get(reverse("api_health")).json()
    assert data["ok"] and data["items"] == 3


def test_security_headers(client):
    response = client.get("/")
    assert "script-src 'self'" in response["Content-Security-Policy"]
    assert response["X-Content-Type-Options"] == "nosniff"


def test_login_wall_when_enabled(client, settings):
    settings.PINKSHEET = {**settings.PINKSHEET, "REQUIRE_LOGIN": True}
    assert client.get("/").status_code == 302
    assert client.get("/api/palette/?recent=1").status_code == 401
    assert client.get("/api/health/").status_code == 200


def test_manage_users_button_only_for_staff(client, django_user_model):
    assert "Manage users" not in client.get("/").content.decode()
    client.force_login(django_user_model.objects.create_user("sam", password="x"))
    assert "Manage users" not in client.get("/").content.decode()
    client.force_login(django_user_model.objects.create_user("lead", password="x", is_staff=True))
    assert 'href="/admin/auth/user/"' in client.get("/").content.decode()
    client.force_login(django_user_model.objects.create_superuser("boss", password="x"))
    assert 'href="/admin/auth/user/"' in client.get("/").content.decode()


def test_sign_out_returns_to_the_app_sign_in(client, settings, django_user_model):
    settings.PINKSHEET = {**settings.PINKSHEET, "REQUIRE_LOGIN": True}
    django_user_model.objects.create_user("sam", password="pw12345678")
    assert client.login(username="sam", password="pw12345678")
    assert client.post("/logout/")["Location"] == "/login/"
    assert client.post("/login/", {"username": "sam", "password": "pw12345678"})["Location"] == "/"


def test_maintenance_mode(client, settings):
    settings.PINKSHEET = {**settings.PINKSHEET, "MAINTENANCE_MODE": True}
    assert client.get("/").status_code == 503
    assert client.get("/api/items/X/").json()["maintenance"] is True


# ── Dispodex branding ────────────────────────────────────────────────────────
@pytest.mark.django_db
@pytest.mark.parametrize("url", ["/", "/login/", "/nope/"])
def test_pages_are_branded_dispodex(client, url):
    html = client.get(url).content.decode()
    assert "Dispodex" in html and "Pinksheet</" not in html
    assert "img/favicon" in html and "manifest" in html and "apple-touch-icon" in html


@pytest.mark.django_db
def test_sidebar_uses_multicolour_logo(client):
    html = client.get("/").content.decode()
    assert 'class="brand-mark" src="/static/img/logo' in html


def test_icon_files_exist_and_manifest_is_valid(settings):
    static = Path(settings.BASE_DIR) / "static"
    manifest = json.loads((static / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["name"] == "Dispodex"
    for icon in manifest["icons"]:
        assert (static / icon["src"]).exists(), icon["src"]
    logo = (static / "img" / "logo.svg").read_text(encoding="utf-8")
    assert len(set(re.findall(r"#[0-9a-fA-F]{6}", logo))) >= 5  # more than just two colours


def test_ebay_wording_comes_from_the_private_file(settings, tmp_path):
    custom = tmp_path / "ebay_boilerplate.txt"
    settings.PINKSHEET = {**settings.PINKSHEET, "EBAY_BOILERPLATE_FILE": custom}
    assert scripts.final_boilerplate() == scripts.DEFAULT_BOILERPLATE  # no file yet: generic example
    custom.write_text("Our policies\nShips Mon-Fri\n", encoding="utf-8")
    assert scripts.build_final_script("Great laptop") == "Our policies\nShips Mon-Fri\n\nGreat laptop"


def test_prints_show_active_inactive(client, stock):
    Item.objects.filter(sku_normalized="LAP-1").update(reviewed=1)
    Item.objects.filter(sku_normalized="LAP-2").update(reviewed=2)  # what saving into the Sold lane sets
    assert b'id="print-review" data-reviewed="1"' in client.get("/intake/?sku=LAP-1").content
    assert b'data-reviewed="0"' in client.get("/intake/?new=1").content
    assert b'<span class="pill">ACTIVE</span>' in client.get("/print/LAP-1/").content
    assert b'<span class="pill">SOLD</span>' in client.get("/print/LAP-2/").content


def test_staff_edit_listing_notes_in_the_browser(client, settings, tmp_path, django_user_model):
    notes_file = tmp_path / "data" / "ebay_boilerplate.txt"
    settings.PINKSHEET = {**settings.PINKSHEET, "EBAY_BOILERPLATE_FILE": notes_file, "REQUIRE_LOGIN": True}
    url = reverse("listing_notes")
    boss = django_user_model.objects.create_user("boss", password="x", first_name="Boss", is_staff=True)
    client.force_login(boss)
    assert "No listing notes are saved on this server yet" in client.get(url).content.decode()

    page = client.post(url, {"text": "Please Read This First  \r\nShips Tue-Sat\r\n\r\n"}, follow=True).content.decode()
    assert "Listing notes saved" in page and "Last changed by Boss" in page
    assert notes_file.read_text(encoding="utf-8") == "Please Read This First\nShips Tue-Sat\n"
    assert scripts.final_boilerplate() == "Please Read This First\nShips Tue-Sat"
    assert "listing notes file" not in client.get("/scripts/").content.decode()
    assert "can&#x27;t be empty" in client.post(url, {"text": "   "}).content.decode()  # refused, file kept
    assert notes_file.read_text(encoding="utf-8") == "Please Read This First\nShips Tue-Sat\n"


def test_only_staff_can_change_listing_notes(client, settings, tmp_path, django_user_model):
    notes_file = tmp_path / "ebay_boilerplate.txt"
    notes_file.write_text("Original notes\n", encoding="utf-8")
    settings.PINKSHEET = {**settings.PINKSHEET, "EBAY_BOILERPLATE_FILE": notes_file, "REQUIRE_LOGIN": True}
    client.force_login(django_user_model.objects.create_user("jo", password="x"))
    html = client.get(reverse("listing_notes")).content.decode()
    assert "Only staff can change" in html and "Original notes" in html and "<textarea" not in html
    client.post(reverse("listing_notes"), {"text": "Hijacked"})
    assert notes_file.read_text(encoding="utf-8") == "Original notes\n"


# ── eBay print sheet ─────────────────────────────────────────────────────────
LOGO = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><rect width="10" height="10"/></svg>'


@pytest.fixture
def logo_dir(settings, tmp_path):
    settings.PINKSHEET = {**settings.PINKSHEET, "PRINT_LOGO_DIR": tmp_path}
    return tmp_path


def upload_logo(client, name, data, content_type="image/svg+xml"):
    from django.core.files.uploadedfile import SimpleUploadedFile

    return client.post(reverse("print_logo_settings"), {"logo": SimpleUploadedFile(name, data, content_type=content_type)})


def test_intake_tools_offer_the_ebay_sheet_with_the_shop_logo(client, logo_dir, stock):
    (logo_dir / "print_logo.svg").write_text(LOGO)
    html = client.get(reverse("intake") + "?sku=LAP-1").content.decode()
    assert 'id="print-ebay-btn"' in html and "Print eBay sheet" in html
    assert f'<div class="print-logo"><img src="{reverse("print_logo")}?v=' in html
    assert 'id="add-print-logo"' not in html
    # The parts the eBay sheet hides are marked so print.css can hide them.
    assert 'class="print-line print-status"' in html and 'class="print-price"' in html
    assert 'class="print-box print-details"' in html and 'class="print-gallery"' in html
    response = client.get(reverse("print_logo"))
    assert response.status_code == 200 and response["Content-Type"] == "image/svg+xml"
    assert "sandbox" in response["Content-Security-Policy"]
    assert response.content.decode() == LOGO
    assert client.get(reverse("print_logo"), HTTP_IF_NONE_MATCH=response["ETag"]).status_code == 304


def test_without_a_logo_the_tools_point_to_where_to_add_it(client, logo_dir, stock):
    html = client.get(reverse("intake") + "?sku=LAP-1").content.decode()
    assert '<div class="print-logo"></div>' in html
    assert f'id="add-print-logo" href="{reverse("print_logo_settings")}"' in html
    assert client.get(reverse("print_logo")).status_code == 404


def test_staff_upload_the_logo_in_the_browser_and_it_prints_right_away(client, logo_dir, stock):
    assert "No logo is saved on this server yet" in client.get(reverse("print_logo_settings")).content.decode()
    response = upload_logo(client, "Shop logo.svg", LOGO.encode())
    assert response.status_code == 302
    assert (logo_dir / "print_logo.svg").read_text() == LOGO
    assert "Current print logo" in client.get(reverse("print_logo_settings")).content.decode()
    assert client.get(reverse("print_logo")).content.decode() == LOGO
    assert f'<div class="print-logo"><img src="{reverse("print_logo")}?v=' in client.get("/intake/?sku=LAP-1").content.decode()

    # A PNG replaces the SVG.
    buffer = io.BytesIO()
    from PIL import Image
    Image.new("RGB", (40, 40), "white").save(buffer, "PNG")
    assert upload_logo(client, "logo.png", buffer.getvalue(), "image/png").status_code == 302
    assert not (logo_dir / "print_logo.svg").exists() and (logo_dir / "print_logo.png").exists()
    assert client.get(reverse("print_logo"))["Content-Type"] == "image/png"

    client.post(reverse("print_logo_settings"), {"action": "remove"})
    assert not list(logo_dir.glob("print_logo.*"))


@pytest.mark.parametrize("data, message", [
    (b"not an image at all", "must be an SVG, PNG or JPG"),
    (b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>', "scripts"),
    (b'<svg xmlns="http://www.w3.org/2000/svg"><rect onload="alert(1)"/></svg>', "scripts"),
    (b'<svg xmlns="http://www.w3.org/2000/svg" xmlns:x="http://www.w3.org/1999/xlink"><image x:href="http://evil.example/a.png"/></svg>', "outside files"),
    (b'<?xml version="1.0"?><!DOCTYPE svg [<!ENTITY a "aaaa">]><svg xmlns="http://www.w3.org/2000/svg">&a;</svg>', "DOCTYPE"),
])
def test_unsafe_or_broken_logos_are_refused(client, logo_dir, data, message):
    response = upload_logo(client, "logo.svg" if data.lstrip().startswith(b"<") else "logo.png", data)
    assert response.status_code == 400 and message in response.content.decode()
    assert not list(logo_dir.glob("print_logo.*"))


def test_only_staff_change_the_logo_when_sign_in_is_on(client, logo_dir, settings, django_user_model):
    settings.PINKSHEET = {**settings.PINKSHEET, "REQUIRE_LOGIN": True}
    client.force_login(django_user_model.objects.create_user("floor", password="x"))
    upload_logo(client, "logo.svg", LOGO.encode())
    assert not list(logo_dir.glob("print_logo.*"))
    assert "Only staff can change the print logo" in client.get(reverse("print_logo_settings")).content.decode()
