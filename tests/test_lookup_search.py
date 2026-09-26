"""Lookup: search every field, several words at once, and the "More filters" panel."""
from datetime import date

import pytest
from django.urls import reverse

from inventory.models import Item
from inventory.services.search import ItemFilters, filter_items

pytestmark = pytest.mark.django_db


@pytest.fixture
def stock(db):
    Item.objects.create(sku="LAP-1", what_is_it="Laptop", brand_model="Dell Latitude 5590", cpu="i5-8350U",
                        ram="16 GB", ssd_gb="256", os="Windows 11", condition="Great", functional="Yes",
                        where_it_goes="Shelf A", source="School district", ready=True, date_received=date(2026, 1, 10),
                        notes="Small scratch on lid", price="120")
    Item.objects.create(sku="LAP-2", what_is_it="Laptop", brand_model="HP EliteBook 840", cpu="i7-8650U",
                        ram="8 GB", ssd_gb="512", condition="Fair", functional="Unknown", where_it_goes="Shelf B",
                        serial_number="5CG1234XYZ", graphics_card="Intel UHD", date_received=date(2026, 3, 2))
    Item.objects.create(sku="DESK-1", what_is_it="Desktop", brand_model="Dell OptiPlex 7050", cpu="i5-7500",
                        ram="16 GB", condition="Scrap", notes="No hard drive")


def skus(query: str) -> set[str]:
    from django.http import QueryDict
    return set(filter_items(ItemFilters.from_query(QueryDict(query))).values_list("sku_normalized", flat=True))


@pytest.mark.parametrize("query,expected", [
    ("q=dell", {"LAP-1", "DESK-1"}),               # brand
    ("q=5590", {"LAP-1"}),                         # model number
    ("q=i7", {"LAP-2"}),                           # CPU
    ("q=5CG1234", {"LAP-2"}),                      # serial
    ("q=scratch", {"LAP-1"}),                      # notes
    ("q=shelf+elitebook", {"LAP-2"}),              # location + model, two words
    ("q=dell+16gb", set()),                        # "16gb" isn't written like that anywhere
    ("q=dell+16", {"LAP-1", "DESK-1"}),            # every word must match somewhere
    ("q=dell+laptop", {"LAP-1"}),
    ("q=desk-1", {"DESK-1"}),                      # SKU still works
    ("q=lap", {"LAP-1", "LAP-2"}),                 # SKU prefix
])
def test_main_box_searches_every_field(stock, query, expected):
    assert skus(query) == expected


@pytest.mark.parametrize("query,expected", [
    ("brand=dell", {"LAP-1", "DESK-1"}),
    ("brand=dell&what=laptop", {"LAP-1"}),
    ("cpu=i5", {"LAP-1", "DESK-1"}),
    ("ram=16", {"LAP-1", "DESK-1"}),
    ("storage=512", {"LAP-2"}),
    ("os=windows", {"LAP-1"}),
    ("gpu=uhd", {"LAP-2"}),
    ("serial=5cg", {"LAP-2"}),
    ("location=shelf", {"LAP-1", "LAP-2"}),
    ("source=school", {"LAP-1"}),
    ("notes=drive", {"DESK-1"}),
    ("condition=fair", {"LAP-2"}),
    ("condition=unicorn", {"LAP-1", "LAP-2", "DESK-1"}),  # not a grade: ignored, not an error
    ("functional=unknown", {"LAP-2"}),
    ("ready=yes", {"LAP-1"}),
    ("ready=no", {"LAP-2", "DESK-1"}),
    ("received_from=2026-02-01", {"LAP-2"}),
    ("received_to=2026-02-01", {"LAP-1"}),
    ("received_from=not-a-date", {"LAP-1", "LAP-2", "DESK-1"}),
    ("q=dell&cpu=i5&condition=scrap", {"DESK-1"}),
])
def test_more_filters(stock, query, expected):
    assert skus(query) == expected


def test_filters_survive_the_round_trip_to_links_and_exports(client, stock):
    from django.http import QueryDict
    filters = ItemFilters.from_query(QueryDict("brand=dell&condition=Great&ready=yes&received_from=2026-01-01"))
    assert filters.as_query() == {"brand": "dell", "condition": "Great", "ready": "yes", "received_from": "2026-01-01"}
    assert filters.advanced_count == 4 and filters.is_filtered
    rows = list(csv_rows(client.get(reverse("export_csv") + "?brand=dell&what=desktop")))
    assert len(rows) == 2 and "DESK-1" in rows[1]


def csv_rows(response):
    import csv
    import io
    return csv.reader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig")
                                  if hasattr(response, "streaming_content") else response.content.decode("utf-8-sig")))


# ── partner exports ─────────────────────────────────────────────────────────
INTERNAL_TEXT = ["School district", "Shelf A", "Shelf B", "Small scratch on lid", "No hard drive"]
INTERNAL_HEADERS = {"Source", "Where It Goes", "Notes", "Status", "What Box", "In eBay Room", "eBay Status",
                    "Date Received", "ID", "Created", "Updated"}


def test_partner_csv_keeps_item_details_but_no_internal_info(client, stock):
    Item.objects.create(sku="SOLD-1", what_is_it="Laptop", brand_model="Dell Latitude 5400", status="sold")
    response = client.get(reverse("export_partner_csv") + "?brand=dell")
    rows = list(csv_rows(response))
    header, body = rows[0], rows[1:]
    assert not INTERNAL_HEADERS & set(header)
    assert {"SKU", "CPU", "RAM", "Storage", "Serial Number", "Condition", "Qty", "Price"} <= set(header)
    assert {row[0] for row in body} == {"LAP-1", "DESK-1"}  # filter applied, sold item left out
    text = response.content.decode("utf-8-sig")
    assert not [word for word in INTERNAL_TEXT if word in text]
    lap = dict(zip(header, next(row for row in body if row[0] == "LAP-1")))
    assert lap["CPU"] == "i5-8350U" and lap["RAM"] == "16 GB" and lap["Price"] == "120.00"
    assert "inventory_partner_" in response["Content-Disposition"]


def test_partner_xlsx_and_zip_leave_out_internal_info(client, stock):
    import io
    import zipfile

    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(b"".join(client.get(reverse("export_partner_xlsx")).streaming_content)))
    sheet = book.active
    headers = [cell.value for cell in sheet[1]]
    assert headers[0] == "SKU" and headers[-1] == "Photos" and not INTERNAL_HEADERS & set(headers)
    values = {str(cell.value) for row in sheet.iter_rows() for cell in row}
    assert not [word for word in INTERNAL_TEXT if word in values]
    archive = zipfile.ZipFile(io.BytesIO(b"".join(client.get(reverse("export_partner_zip")).streaming_content)))
    info = archive.read("LAP-1/info.txt").decode()
    assert "CPU: i5-8350U" in info and not [word for word in INTERNAL_TEXT if word in info]


def test_partner_sortable_xlsx_has_filters_numbers_and_no_photos(client, stock, image):
    import io

    import openpyxl
    client.post(reverse("api_photo_upload"), {"sku": "LAP-1", "photo": image()})
    Item.objects.create(sku="SOLD-2", what_is_it="Laptop", status="sold")
    response = client.get(reverse("export_partner_sortable_xlsx"))
    assert "inventory_partner_sortable_" in response["Content-Disposition"]
    sheet = openpyxl.load_workbook(io.BytesIO(b"".join(response.streaming_content))).active
    headers = [cell.value for cell in sheet[1]]
    assert headers[0] == "SKU" and "Photos" not in headers and not INTERNAL_HEADERS & set(headers)
    assert sheet.auto_filter.ref.startswith("A1:") and not sheet._images
    rows = {row[0]: dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert set(rows) == {"LAP-1", "LAP-2", "DESK-1"}
    assert rows["LAP-1"]["Price"] == 120 and rows["LAP-1"]["Qty"] == 1  # real numbers, so they sort properly
    values = {str(v) for row in rows.values() for v in row.values()}
    assert not [word for word in INTERNAL_TEXT if word in values]


def test_regular_export_still_has_every_column(client, stock):
    header = next(csv_rows(client.get(reverse("export_csv"))))
    assert INTERNAL_HEADERS <= set(header) and header[0] == "SKU"


def test_lookup_page_shows_panel_chips_and_suggestions(client, stock):
    html = client.get(reverse("lookup") + "?brand=dell&cpu=i5").content.decode()
    assert 'id="more-filters"' in html and 'id="more-filters" hidden' not in html  # open while in use
    assert "Brand &amp; model:</span> dell" in html and "CPU:</span> i5" in html
    assert 'value="Dell Latitude 5590"' in html  # brand suggestions from existing items
    assert "LAP-1" in html and "DESK-1" in html and "LAP-2" not in html
    # Removing one chip keeps the other filter.
    assert "?cpu=i5" in html and "?brand=dell" in html


def test_panel_is_closed_when_unused(client, stock):
    html = client.get(reverse("lookup")).content.decode()
    assert 'id="more-filters" hidden' in html and "filter-chips" not in html


def test_old_links_with_gap_and_stale_still_work(client, stock):
    Item.objects.filter(sku_normalized="LAP-1").update(price=None)
    html = client.get(reverse("lookup") + "?gap=no-price").content.decode()
    assert "LAP-1" in html and "Missing:</span> price" in html


def test_board_cards_include_brand_and_condition(client, stock):
    cards = {c["sku"]: c for c in client.get(reverse("api_board_cards")).json()["lanes"]["intake"]}
    assert cards["LAP-1"]["brand"] == "Dell Latitude 5590" and cards["LAP-1"]["condition"] == "Great"
