"""Downloads: inventory CSV, ZIP bundle with photo folders, and an Excel sheet with photos."""
from __future__ import annotations

import csv
import io
import tempfile
import threading
import zipfile
from collections import defaultdict
from pathlib import Path

from django.utils import timezone

from core.skus import sanitize_filename
from inventory.models import Item, Photo, Status
from inventory.services import photos as photo_service

# (field, header) in the order columns appear in every export.
EXPORT_COLUMNS = [
    ("sku", "SKU"),
    ("status", "Status"),
    ("what_is_it", "What is it?"),
    ("brand_model", "Brand/Model"),
    ("source", "Source"),
    ("date_received", "Date Received"),
    ("condition", "Condition"),
    ("functional", "Functional"),
    ("power_on", "Power On"),
    ("price", "Price"),
    ("quantity", "Qty"),
    ("ebay_category", "eBay Category"),
    ("ebay_status", "eBay Status"),
    ("where_it_goes", "Where It Goes"),
    ("cpu", "CPU"),
    ("ram", "RAM"),
    ("ssd_gb", "SSD (GB)"),
    ("graphics_card", "Graphics Card"),
    ("screen_resolution", "Screen Resolution"),
    ("battery_health", "Battery Health"),
    ("os", "OS"),
    ("compatible_os", "Compatible OS"),
    ("wifi_card_installed", "WiFi Card Installed"),
    ("picture_taken", "Picture Taken"),
    ("cords_adapters", "Cords/Adapters"),
    ("keep_items_together", "Keep Items Together"),
    ("what_box", "What Box"),
    ("in_ebay_room", "In eBay Room"),
    ("diagnostics_test_ran", "Diagnostics Ran"),
    ("serial_number", "Serial Number"),
    ("notes", "Notes"),
    ("ebay_category_path", "eBay Category Path"),
    ("ebay_category_id", "eBay Category ID"),
    ("id", "ID"),
    ("created_at", "Created"),
    ("updated_at", "Updated"),
]

# What a partner may see: everything that describes the item itself. Left out on purpose:
# where it came from (source), where it sits (location, box, eBay room), notes, the
# workflow lane, eBay bookkeeping and internal IDs/dates.
PARTNER_COLUMNS = [
    ("sku", "SKU"),
    ("what_is_it", "Item"),
    ("brand_model", "Brand/Model"),
    ("ebay_category", "Category"),
    ("condition", "Condition"),
    ("functional", "Functional"),
    ("power_on", "Powers On"),
    ("cpu", "CPU"),
    ("ram", "RAM"),
    ("ssd_gb", "Storage"),
    ("graphics_card", "Graphics"),
    ("screen_resolution", "Screen Resolution"),
    ("battery_health", "Battery Health"),
    ("os", "OS"),
    ("compatible_os", "Compatible OS"),
    ("wifi_card_installed", "WiFi Card Installed"),
    ("diagnostics_test_ran", "Diagnostics Ran"),
    ("cords_adapters", "Cords/Adapters Included"),
    ("serial_number", "Serial Number"),
    ("quantity", "Qty"),
    ("price", "Price"),
]

# Only one big (photo) export at a time — they can take a while on a busy day.
heavy_export_lock = threading.Lock()


def cell_value(item: Item, field: str, for_sheet: bool = False):
    value = getattr(item, field)
    if field == "status":
        return item.status_label
    if field == "ebay_category":
        return item.ebay_category or item.ebay_category_path
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if value is None:
        return ""
    if field in ("created_at", "updated_at"):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M:%S")
    if field == "date_received":
        return value.isoformat()
    if field == "price":
        return float(value) if for_sheet else f"{value:.2f}"
    if field == "quantity" and for_sheet:
        return int(value)
    return value


def filename(prefix: str, scope: str, ext: str) -> str:
    suffix = "active_" if scope == "active" else ""
    return f"{prefix}_{suffix}{timezone.localdate():%Y-%m-%d}.{ext}"


def inventory_csv(items, columns=EXPORT_COLUMNS) -> str:
    buffer = io.StringIO()
    buffer.write("﻿")  # BOM so Excel opens UTF-8 correctly
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow([header for _, header in columns])
    for item in items:
        writer.writerow([cell_value(item, field) for field, _ in columns])
    return buffer.getvalue()


def _photos_by_sku(items: list[Item]) -> dict[str, list[Photo]]:
    grouped: dict[str, list[Photo]] = defaultdict(list)
    for photo in Photo.objects.filter(sku_normalized__in={i.sku_normalized for i in items}).order_by("sort_order", "id"):
        grouped[photo.sku_normalized].append(photo)
    return grouped


def _folder_name(item: Item, used: set[str]) -> str:
    base = sanitize_filename(item.sku or item.sku_normalized, fallback=f"UNASSIGNED-{item.pk}")
    candidate, n = base, 2
    while candidate.lower() in used:
        candidate = f"{base}-{n}"
        n += 1
    used.add(candidate.lower())
    return candidate


def _info_text(item: Item, photo_count: int, columns) -> str:
    lines = [f"Photos: {photo_count}"]
    lines += [f"{header}: {cell_value(item, field)}" for field, header in columns]
    return "\n".join(lines) + "\n"


def inventory_zip(items: list[Item], csv_name: str, columns=EXPORT_COLUMNS):
    """A ZIP with one folder per SKU (photos + info.txt) and the CSV at the root.

    Returns an open temporary file positioned at the start; it deletes itself when closed.
    """
    handle = tempfile.TemporaryFile()
    photos = _photos_by_sku(items)
    used: set[str] = set()
    with zipfile.ZipFile(handle, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(csv_name, inventory_csv(items, columns))
        for item in items:
            folder = _folder_name(item, used)
            archive.writestr(f"{folder}/", "")
            count = 0
            for photo in photos.get(item.sku_normalized, []):
                try:
                    photo_service.widen_photo(photo)  # in case start-up hasn't reached it yet
                except (OSError, ValueError, SyntaxError):
                    pass
                path = photo_service.photo_path(photo)
                if not path.exists():
                    continue
                count += 1
                base = sanitize_filename(Path(photo.original_name).stem)
                # Photos are already compressed; storing them avoids wasted CPU.
                archive.write(path, f"{folder}/{count:02d}_{base}{path.suffix}", compress_type=zipfile.ZIP_STORED)
            archive.writestr(f"{folder}/info.txt", _info_text(item, count, columns))
    handle.seek(0)
    return handle


def inventory_xlsx(items: list[Item], columns=EXPORT_COLUMNS):
    """An Excel workbook, one row per SKU, with every photo shown in the last column."""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SheetImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.units import pixels_to_EMU

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventory"
    fields = [field for field, _ in columns]
    headers = [header for _, header in columns] + ["Photos"]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="E7F0FA")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "B2"

    photo_col = len(headers)
    photo_letter = get_column_letter(photo_col)
    photos = _photos_by_sku(items)
    thumb_height = 96
    for row_index, item in enumerate(items, start=2):
        sheet.append([cell_value(item, field, for_sheet=True) for field in fields] + [""])
        if "price" in fields:
            sheet.cell(row=row_index, column=fields.index("price") + 1).number_format = '"$"#,##0.00'
        total_height, placed = 0, 0
        for photo in photos.get(item.sku_normalized, []):
            thumb = photo_service.thumbnail_file(photo, size=160)
            if not thumb:
                continue
            image = SheetImage(str(thumb))
            scale = thumb_height / image.height if image.height else 1
            image.width, image.height = int(image.width * scale), thumb_height
            # Stack the photos down the cell with a small gap between them.
            marker = AnchorMarker(col=photo_col - 1, colOff=pixels_to_EMU(4), row=row_index - 1,
                                  rowOff=pixels_to_EMU(4 + total_height))
            size = XDRPositiveSize2D(pixels_to_EMU(image.width), pixels_to_EMU(image.height))
            image.anchor = OneCellAnchor(_from=marker, ext=size)
            sheet.add_image(image)
            total_height += thumb_height + 4
            placed += 1
        if placed:
            sheet.row_dimensions[row_index].height = round((total_height + 8) * 0.75, 1)
        for cell in sheet[row_index]:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column_letter != photo_letter)

    widths = {"sku": 18, "what_is_it": 28, "brand_model": 28, "notes": 40, "ebay_category_path": 40,
              "price": 12, "quantity": 6, "created_at": 19, "updated_at": 19}
    for index, field in enumerate(fields, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(field, 16)
    sheet.column_dimensions[photo_letter].width = 26

    handle = tempfile.TemporaryFile()
    workbook.save(handle)
    handle.seek(0)
    return handle


def sortable_xlsx(items: list[Item], columns):
    """A plain Excel sheet without photos, with filter buttons on, so it can be sorted and filtered safely.

    (Photos float over the cells in the photo workbook, so sorting there would mix them up.)
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventory"
    fields = [field for field, _ in columns]
    sheet.append([header for _, header in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E7F0FA")
    for item in items:
        sheet.append([cell_value(item, field, for_sheet=True) for field in fields])
    if "price" in fields:
        letter = get_column_letter(fields.index("price") + 1)
        for cell in sheet[letter][1:]:
            cell.number_format = '"$"#,##0.00'
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{max(sheet.max_row, 2)}"
    widths = {"sku": 18, "what_is_it": 28, "brand_model": 28, "ebay_category": 24, "price": 12, "quantity": 6}
    for index, field in enumerate(fields, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(field, 16)

    handle = tempfile.TemporaryFile()
    workbook.save(handle)
    handle.seek(0)
    return handle


def active_scope(items_qs, scope: str):
    return items_qs.exclude(status=Status.SOLD) if scope == "active" else items_qs
